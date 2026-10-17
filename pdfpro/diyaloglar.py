# -*- coding: utf-8 -*-
"""
PDF PRO - Ayar pencereleri (imza, antet, filigran, sayfa no, bol, birlestir, sifre, OCR, metin)
Her pencere 'Uygula' denince sonuc sozlugunu self.sonuc'a koyar ve kapanir.
"""
import os
import sys
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

from . import cekirdek as C

WINDOWS = sys.platform.startswith("win")
FONT = "Segoe UI" if WINDOWS else "DejaVu Sans"
R_KART = "#FFFFFF"
R_BASLIK = "#312E81"
R_SOLUK = "#64748B"
R_VURGU = "#4F46E5"

RESIM_TUR = [("Resim", "*.png *.jpg *.jpeg *.bmp *.tif *.tiff *.gif *.webp"), ("Tüm dosyalar", "*.*")]
PDF_RESIM_TUR = [("PDF veya resim", "*.pdf *.png *.jpg *.jpeg *.bmp *.tif *.tiff"), ("Tüm dosyalar", "*.*")]


def baslik_rengi(ust):
    """Pencere basligi, acik menunun rengini alir (Dosya mavi, Dönüştür yesil ...)."""
    try:
        from .arayuz import MENU_RENKLERI
        return MENU_RENKLERI.get(getattr(ust, "_secili_menu", None), R_VURGU)
    except Exception:
        return R_VURGU


class Diyalog(tk.Toplevel):
    def __init__(self, ust, baslik, aciklama=None, dugme="Uygula", genislik=440):
        super().__init__(ust)
        self.u = ust
        self.sonuc = None
        self.title(baslik)
        self.configure(bg=R_KART)
        self.transient(ust)
        self.resizable(False, False)
        renk = baslik_rengi(ust)
        ust_bant = tk.Frame(self, bg=renk)
        ust_bant.pack(fill="x")
        tk.Label(ust_bant, text=baslik, bg=renk, fg="white", font=(FONT, 13, "bold")).pack(anchor="w", padx=18, pady=(12, 10))
        if aciklama:
            tk.Label(self, text=aciklama, bg=R_KART, fg=R_SOLUK, font=(FONT, 9), wraplength=genislik - 30,
                     justify="left").pack(anchor="w", padx=18, pady=(10, 0))
        self.govde = ttk.Frame(self, style="Kart.TFrame", padding=(18, 10))
        self.govde.pack(fill="both", expand=True)
        self.govde.columnconfigure(1, weight=1)
        alt = ttk.Frame(self, style="Kart.TFrame", padding=(18, 4, 18, 14))
        alt.pack(fill="x")
        ttk.Button(alt, text="Vazgeç", style="Ikincil.TButton", command=self.destroy).pack(side="right")
        self.ana_dugme = ttk.Button(alt, text=dugme, style="Ana.TButton", command=self._tamam)
        self.ana_dugme.pack(side="right", padx=(0, 8))
        self.bind("<Escape>", lambda e: self.destroy())
        self.minsize(genislik, 10)
        self._satir = 0

    def ac(self):
        self.update_idletasks()
        u = self.u
        x = u.winfo_rootx() + (u.winfo_width() - self.winfo_width()) // 2
        y = u.winfo_rooty() + (u.winfo_height() - self.winfo_height()) // 3
        self.geometry(f"+{max(0, x)}+{max(0, y)}")
        self.grab_set()
        self.focus_set()
        self.wait_window()
        return self.sonuc

    def _tamam(self):
        try:
            s = self.topla()
        except ValueError as e:
            messagebox.showwarning("PDF PRO", str(e), parent=self)
            return
        if s is None:
            return
        self.sonuc = s
        self.destroy()

    def topla(self):
        return {}

    # -- yardimcilar
    def satir(self, etiket, widget=None, ipucu=None):
        r = self._satir
        if etiket:
            ttk.Label(self.govde, text=etiket, style="Kart.TLabel").grid(row=r, column=0, sticky="w", pady=4, padx=(0, 10))
        if widget is not None:
            widget.grid(row=r, column=1, sticky="ew", pady=4)
        self._satir += 1
        if ipucu:
            ttk.Label(self.govde, text=ipucu, style="Soluk.TLabel").grid(row=self._satir, column=1, sticky="w")
            self._satir += 1
        return widget

    def tam_satir(self, widget, pady=4):
        widget.grid(row=self._satir, column=0, columnspan=2, sticky="w", pady=pady)
        self._satir += 1
        return widget

    def dosya_alani(self, etiket, var, turler):
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        e = ttk.Entry(f, textvariable=var, width=34)
        e.pack(side="left", fill="x", expand=True)

        def sec():
            y = filedialog.askopenfilename(parent=self, filetypes=turler,
                                           initialdir=os.path.dirname(var.get()) if var.get() else None)
            if y:
                var.set(os.path.normpath(y))

        ttk.Button(f, text="📂", width=3, style="Ikincil.TButton", command=sec).pack(side="left", padx=(4, 0))
        self.satir(etiket, f)
        return e

    def sayfa_alani(self, varsayilan=""):
        e = ttk.Entry(self.govde)
        e.insert(0, varsayilan)
        self.satir("Sayfalar:", e, "boş = tümü · son · 1-3, 5 · tek · çift")
        return e

    def kaydirici(self, etiket, nesne_var, bas, son, bicim):
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        et = ttk.Label(f, text=bicim(nesne_var.get()), width=7, style="Kart.TLabel")
        ttk.Scale(f, from_=bas, to=son, variable=nesne_var,
                  command=lambda v: et.config(text=bicim(float(v)))).pack(side="left", fill="x", expand=True)
        et.pack(side="left")
        self.satir(etiket, f)


# ================================================================== imza / kase
class ImzaDiyalog(Diyalog):
    def __init__(self, ust, ayar):
        super().__init__(ust, "✍ İmza / Kaşe Yerleştir",
                         "Resmi seçin, sonra sayfada yerleştirmek istediğiniz yere TIKLAYIN "
                         "(ya da fareyle bir kutu çizin, imza o kutuya sığdırılır).",
                         dugme="Sayfaya Yerleştir →", genislik=470)
        self.ayar = ayar
        self.gecmis = [y for y in ayar.get("imza_gecmis", []) if os.path.isfile(y)]
        self.yol = tk.StringVar(value=self.gecmis[0] if self.gecmis else "")
        self._fotolar = []
        if self.gecmis:
            ttk.Label(self.govde, text="Son kullanılanlar — tıklayın, hemen sayfaya koyun:", style="Kart.TLabel").grid(row=self._satir, column=0, columnspan=2, sticky="w")
            self._satir += 1
            f = ttk.Frame(self.govde, style="Kart.TFrame")
            self.tam_satir(f)
            self._gecmis_cerceve = f
            self._gecmis_ciz()
        self.dosya_alani("Resim:", self.yol, RESIM_TUR)
        self.gen = tk.DoubleVar(value=ayar.get("imza_genislik", 45))
        self.kaydirici("Genişlik:", self.gen, 15, 160, lambda v: f"{float(v):.0f} mm")
        self.beyaz = tk.BooleanVar(value=ayar.get("imza_beyaz", True))
        self.tam_satir(ttk.Checkbutton(self.govde, text="Beyaz arka planı kaldır (taranmış imza / kaşe için)",
                                       variable=self.beyaz, style="Kart.TCheckbutton"))
        self.tum = tk.BooleanVar(value=False)
        self.tam_satir(ttk.Checkbutton(self.govde, text="Tüm sayfalarda aynı yere koy",
                                       variable=self.tum, style="Kart.TCheckbutton"))

    def _gecmis_ciz(self):
        from PIL import Image, ImageTk
        for w in self._gecmis_cerceve.winfo_children():
            w.destroy()
        self._fotolar = []
        for y in self.gecmis[:6]:
            try:
                im = Image.open(y)
                im.thumbnail((84, 50))
                ft = ImageTk.PhotoImage(im)
            except Exception:
                continue
            self._fotolar.append(ft)
            secili = y == self.yol.get()
            b = tk.Button(self._gecmis_cerceve, image=ft, bg="#E5E7EB" if secili else "white", relief="solid",
                          bd=2 if secili else 1, width=90, height=56, cursor="hand2",
                          command=lambda yy=y: (self.yol.set(yy), self._tamam()))
            b.pack(side="left", padx=3, pady=4)

    def topla(self):
        y = self.yol.get().strip()
        if not os.path.isfile(y):
            raise ValueError("İmza / kaşe resmini seçin (PNG veya JPG).")
        g = [y] + [x for x in self.gecmis if x != y]
        self.ayar["imza_gecmis"] = g[:8]
        self.ayar["imza_genislik"] = round(self.gen.get())
        self.ayar["imza_beyaz"] = self.beyaz.get()
        return {"resim": y, "genislik_mm": self.gen.get(), "beyaz": self.beyaz.get(), "tum": self.tum.get()}


# ================================================================== antet
class AntetDiyalog(Diyalog):
    def __init__(self, ust, ayar):
        super().__init__(ust, "🏢 Antet Ekle",
                         "En iyisi: Word'deki antetli boş kağıdınızı PDF olarak kaydedip burada seçin. "
                         "PNG/JPG antet de olur.")
        self.ayar = ayar
        self.yol = tk.StringVar(value=ayar.get("antet_yolu", ""))
        self.dosya_alani("Antet dosyası:", self.yol, PDF_RESIM_TUR)
        self.sayfa = self.sayfa_alani()
        self.yer = ttk.Combobox(self.govde, values=["Tam sayfa", "Üst bilgi", "Alt bilgi"], state="readonly")
        self.yer.set(ayar.get("antet_yer", "Tam sayfa"))
        self.satir("Resim yerleşimi:", self.yer, "(sadece PNG/JPG antet için)")
        self.onde = tk.BooleanVar(value=False)
        self.tam_satir(ttk.Checkbutton(self.govde, text="Antedi yazıların ÖNÜNE bas (taranmış belgede görünmüyorsa)",
                                       variable=self.onde, style="Kart.TCheckbutton"))
        self.devam = tk.BooleanVar(value=True)
        self.tam_satir(ttk.Checkbutton(self.govde, text="Antet PDF'i 2 sayfalıysa: 1. sayfaya ilkini, diğerlerine ikincisini bas",
                                       variable=self.devam, style="Kart.TCheckbutton"))

    def topla(self):
        y = self.yol.get().strip()
        if not os.path.isfile(y):
            raise ValueError("Antet dosyasını seçin (PDF veya resim).")
        self.ayar["antet_yolu"] = y
        self.ayar["antet_yer"] = self.yer.get()
        return {"antet": y, "sayfalar": self.sayfa.get(), "yerlesim": self.yer.get(),
                "onde": self.onde.get(), "devam": self.devam.get()}


# ================================================================== filigran
class FiligranDiyalog(Diyalog):
    RENKLER = {"Kırmızı": "#C0392B", "Lacivert": "#1F3A5F", "Gri": "#7F8C8D", "Siyah": "#000000", "Yeşil": "#1E8449"}

    def __init__(self, ust, ayar):
        super().__init__(ust, "💧 Filigran Ekle")
        self.ayar = ayar
        self.tur = tk.StringVar(value="metin")
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        ttk.Radiobutton(f, text="Yazı", value="metin", variable=self.tur, style="Kart.TRadiobutton").pack(side="left")
        ttk.Radiobutton(f, text="Resim / Logo", value="resim", variable=self.tur, style="Kart.TRadiobutton").pack(side="left", padx=12)
        self.satir("Tür:", f)
        self.metin = ttk.Combobox(self.govde, values=["GİZLİDİR", "TASLAK", "ASLI GİBİDİR", "KOPYA", "ÖRNEKTİR", "İPTAL"])
        self.metin.set(ayar.get("filigran_metin", "GİZLİDİR"))
        self.satir("Yazı:", self.metin)
        self.resim = tk.StringVar(value=ayar.get("filigran_resim", ""))
        self.dosya_alani("Resim:", self.resim, RESIM_TUR)
        f2 = ttk.Frame(self.govde, style="Kart.TFrame")
        self.renk = ttk.Combobox(f2, values=list(self.RENKLER), state="readonly", width=10)
        self.renk.set(ayar.get("filigran_renk", "Kırmızı"))
        self.renk.pack(side="left")
        ttk.Label(f2, text="   Açı:", style="Kart.TLabel").pack(side="left")
        self.aci = ttk.Spinbox(f2, from_=-90, to=90, increment=15, width=5)
        self.aci.set(45)
        self.aci.pack(side="left")
        self.satir("Renk:", f2)
        self.opak = tk.DoubleVar(value=ayar.get("filigran_opak", 15))
        self.kaydirici("Koyuluk:", self.opak, 5, 60, lambda v: f"%{float(v):.0f}")
        self.desen = tk.BooleanVar(value=False)
        self.tam_satir(ttk.Checkbutton(self.govde, text="Tüm sayfaya döşe (tekrarlı desen)", variable=self.desen,
                                       style="Kart.TCheckbutton"))
        self.sayfa = self.sayfa_alani()

    def topla(self):
        s = {"tur": self.tur.get(), "sayfalar": self.sayfa.get(), "opaklik": self.opak.get() / 100}
        if s["tur"] == "resim":
            if not os.path.isfile(self.resim.get()):
                raise ValueError("Filigran resmini seçin.")
            s["resim"] = self.resim.get()
            self.ayar["filigran_resim"] = s["resim"]
        else:
            if not self.metin.get().strip():
                raise ValueError("Filigran yazısını girin.")
            try:
                s["aci"] = float(self.aci.get())
            except ValueError:
                s["aci"] = 45
            s.update(metin=self.metin.get(), renk=self.RENKLER.get(self.renk.get(), "#C0392B"), desen=self.desen.get())
        self.ayar.update(filigran_metin=self.metin.get(), filigran_renk=self.renk.get(),
                         filigran_opak=round(self.opak.get()))
        return s


# ================================================================== sayfa no
class SayfaNoDiyalog(Diyalog):
    def __init__(self, ust):
        super().__init__(ust, "#️⃣ Sayfa Numarası Ekle", "{n} = sayfa numarası, {N} = toplam sayfa")
        self.bicim = ttk.Combobox(self.govde, values=["Sayfa {n} / {N}", "{n} / {N}", "{n}", "- {n} -", "Page {n} of {N}"])
        self.bicim.set("Sayfa {n} / {N}")
        self.satir("Biçim:", self.bicim)
        self.konum = ttk.Combobox(self.govde, values=["Orta alt", "Sağ alt", "Sol alt", "Orta üst", "Sağ üst", "Sol üst"], state="readonly")
        self.konum.set("Orta alt")
        self.satir("Konum:", self.konum)
        self.boyut = ttk.Spinbox(self.govde, from_=6, to=24, width=6)
        self.boyut.set(10)
        self.satir("Yazı boyu:", self.boyut)
        self.bas = ttk.Spinbox(self.govde, from_=0, to=9999, width=6)
        self.bas.set(1)
        self.satir("Başlangıç no:", self.bas)
        self.sayfa = self.sayfa_alani()

    def topla(self):
        try:
            boyut, bas = float(self.boyut.get()), int(self.bas.get())
        except ValueError:
            raise ValueError("Yazı boyu ve başlangıç numarası sayı olmalı.")
        return {"bicim": self.bicim.get(), "konum": self.konum.get(), "boyut": boyut, "baslangic": bas,
                "sayfalar": self.sayfa.get()}


# ================================================================== bol
class BolDiyalog(Diyalog):
    def __init__(self, ust, sayfa_sayisi):
        super().__init__(ust, "✂ Belgeyi Böl", f"Belgede {sayfa_sayisi} sayfa var. Parçalar yeni bir klasöre kaydedilir.",
                         dugme="Böl")
        self.mod = tk.StringVar(value="her_sayfa")
        self.tam_satir(ttk.Radiobutton(self.govde, text="Her sayfa ayrı PDF", value="her_sayfa",
                                       variable=self.mod, style="Kart.TRadiobutton"))
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        ttk.Radiobutton(f, text="Her", value="n_sayfa", variable=self.mod, style="Kart.TRadiobutton").pack(side="left")
        self.n = ttk.Spinbox(f, from_=1, to=999, width=5)
        self.n.set(2)
        self.n.pack(side="left", padx=4)
        ttk.Label(f, text="sayfada bir", style="Kart.TLabel").pack(side="left")
        self.tam_satir(f)
        f2 = ttk.Frame(self.govde, style="Kart.TFrame")
        ttk.Radiobutton(f2, text="Aralıklara göre:", value="araliklar", variable=self.mod, style="Kart.TRadiobutton").pack(side="left")
        self.aralik = ttk.Entry(f2, width=22)
        self.aralik.insert(0, "1-3, 4-son")
        self.aralik.pack(side="left", padx=4)
        self.aralik.bind("<FocusIn>", lambda e: self.mod.set("araliklar"))
        self.tam_satir(f2)

    def topla(self):
        m = self.mod.get()
        return {"mod": m, "deger": self.n.get() if m == "n_sayfa" else self.aralik.get()}


# ================================================================== birlestir
class BirlestirDiyalog(Diyalog):
    def __init__(self, ust, ayar, ilk=None):
        super().__init__(ust, "🗂 Dosyaları Birleştir",
                         "PDF ve resimleri ekleyin, sırayı ayarlayın. Sonuç yeni bir sekmede açılır.",
                         dugme="Birleştir", genislik=520)
        self.ayar = ayar
        self.yollar = list(ilk or [])
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        f.grid(row=0, column=0, columnspan=2, sticky="nsew")
        self.liste = tk.Listbox(f, selectmode="extended", height=12, width=60, activestyle="none", relief="solid",
                                bd=1, font=(FONT, 10), selectbackground=R_VURGU)
        self.liste.pack(side="left", fill="both", expand=True)
        d = ttk.Frame(f, style="Kart.TFrame")
        d.pack(side="left", fill="y", padx=(8, 0))
        for t, k in (("➕ Ekle", self._ekle), ("▲ Yukarı", lambda: self._tasi(-1)), ("▼ Aşağı", lambda: self._tasi(1)),
                     ("A→Z", self._sirala), ("✖ Kaldır", self._kaldir)):
            ttk.Button(d, text=t, style="Ikincil.TButton", command=k, width=10).pack(fill="x", pady=2)
        try:
            from tkinterdnd2 import DND_FILES
            self.liste.drop_target_register(DND_FILES)
            self.liste.dnd_bind("<<Drop>>", lambda e: self._ekle_yollar(self.tk.splitlist(e.data)))
        except Exception:
            pass
        self._yenile()

    def _ekle(self):
        ys = filedialog.askopenfilenames(parent=self, filetypes=PDF_RESIM_TUR, initialdir=self.ayar.get("son_klasor"))
        self._ekle_yollar(ys)

    def _ekle_yollar(self, ys):
        for y in ys:
            if os.path.isdir(y):
                for ad in sorted(os.listdir(y)):
                    if ad.lower().endswith((".pdf",) + C.RESIM_UZANTILARI):
                        self.yollar.append(os.path.join(y, ad))
            elif os.path.isfile(y):
                self.yollar.append(os.path.normpath(y))
                self.ayar["son_klasor"] = os.path.dirname(y)
        self._yenile()

    def _yenile(self, sec=()):
        self.liste.delete(0, "end")
        for i, y in enumerate(self.yollar, 1):
            self.liste.insert("end", f" {i:>2}.  {os.path.basename(y)}")
        for s in sec:
            self.liste.selection_set(s)

    def _tasi(self, yon):
        sec = list(self.liste.curselection())
        if not sec or (yon < 0 and sec[0] == 0) or (yon > 0 and sec[-1] == len(self.yollar) - 1):
            return
        for i in (sec if yon < 0 else reversed(sec)):
            self.yollar[i], self.yollar[i + yon] = self.yollar[i + yon], self.yollar[i]
        self._yenile([i + yon for i in sec])

    def _sirala(self):
        self.yollar.sort(key=lambda y: os.path.basename(y).lower())
        self._yenile()

    def _kaldir(self):
        for i in reversed(self.liste.curselection()):
            del self.yollar[i]
        self._yenile()

    def topla(self):
        if len(self.yollar) < 1:
            raise ValueError("Listeye en az bir dosya ekleyin.")
        return {"yollar": list(self.yollar)}


# ================================================================== sifre
class SifreDiyalog(Diyalog):
    def __init__(self, ust):
        super().__init__(ust, "🔒 Şifre Koy", "Şifreli bir kopya kaydedilir. Şifreyi unutursanız belge açılamaz!",
                         dugme="Şifreli Kopya Kaydet")
        self.s1 = ttk.Entry(self.govde, show="●")
        self.satir("Şifre:", self.s1)
        self.s2 = ttk.Entry(self.govde, show="●")
        self.satir("Tekrar:", self.s2)
        self.s1.focus_set()

    def topla(self):
        if not self.s1.get():
            raise ValueError("Şifre girin.")
        if self.s1.get() != self.s2.get():
            raise ValueError("Şifreler aynı değil.")
        return {"sifre": self.s1.get()}


class SifreSorDiyalog(Diyalog):
    def __init__(self, ust, ad, deneme):
        super().__init__(ust, "🔒 Şifreli belge",
                         (f"'{ad}' şifreli. Açmak için şifreyi girin." if not deneme else "Şifre yanlış, tekrar deneyin."),
                         dugme="Aç")
        self.s = ttk.Entry(self.govde, show="●", width=30)
        self.satir("Şifre:", self.s)
        self.s.focus_set()
        self.s.bind("<Return>", lambda e: self._tamam())

    def topla(self):
        return {"sifre": self.s.get()}


# ================================================================== sikistir
class SikistirDiyalog(Diyalog):
    def __init__(self, ust, mevcut_boyut=None):
        bilgi = "Küçültülmüş bir kopya kaydedilir; açık belge değişmez."
        if mevcut_boyut:
            bilgi = f"Şu anki boyut: {mevcut_boyut}.  " + bilgi
        super().__init__(ust, "🗜 PDF Sıkıştır", bilgi, dugme="Sıkıştır", genislik=470)
        self.seviye = tk.StringVar(value="orta")
        for d, t in (("hafif", "Hafif — kalite çok iyi (200 dpi)"),
                     ("orta", "Orta — ÖNERİLEN, e-posta ve sistem yüklemeleri için (150 dpi)"),
                     ("guclu", "Güçlü — taramalarda çok küçülür (110 dpi)"),
                     ("en_kucuk", "En küçük — okunur ama kalite düşer (90 dpi)")):
            self.tam_satir(ttk.Radiobutton(self.govde, text=t, value=d, variable=self.seviye,
                                           style="Kart.TRadiobutton"), pady=2)
        self.gri = tk.BooleanVar(value=False)
        self.tam_satir(ttk.Checkbutton(self.govde, text="Siyah-beyaz yap (gri tonlama) — daha da küçülür",
                                       variable=self.gri, style="Kart.TCheckbutton"), pady=(8, 2))
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        self.hedef = ttk.Combobox(f, values=["", "0.5", "1", "2", "3", "5", "10"], width=6)
        self.hedef.pack(side="left")
        ttk.Label(f, text=" MB'ın altına indir (boş = seçilen seviye)", style="Kart.TLabel").pack(side="left")
        self.satir("Hedef boyut:", f, "Örn. sistem 2 MB'tan büyük dosya kabul etmiyorsa 2 yazın; gerekirse seviyeyi kendisi artırır.")

    def topla(self):
        h = self.hedef.get().strip().replace(",", ".")
        try:
            hedef_kb = float(h) * 1024 if h else None
        except ValueError:
            raise ValueError("Hedef boyut sayı olmalı (örn. 2).")
        return {"seviye": self.seviye.get(), "gri": self.gri.get(), "hedef_kb": hedef_kb}


# ================================================================== OCR
class OcrDiyalog(Diyalog):
    def __init__(self, ust, baslik="🔍 Aranabilir PDF (OCR)", aciklama=None, dugme="Başlat"):
        super().__init__(ust, baslik, aciklama or
                         "Taranmış sayfalardaki yazılar okunur; görünüm aynı kalır, metin aranabilir ve "
                         "düzeltilebilir hale gelir.", dugme=dugme)
        f = ttk.Frame(self.govde, style="Kart.TFrame")
        self.tur = tk.BooleanVar(value=True)
        self.eng = tk.BooleanVar(value=True)
        ttk.Checkbutton(f, text="Türkçe", variable=self.tur, style="Kart.TCheckbutton").pack(side="left")
        ttk.Checkbutton(f, text="İngilizce", variable=self.eng, style="Kart.TCheckbutton").pack(side="left", padx=10)
        self.satir("Dil:", f)
        self.dpi = ttk.Combobox(self.govde, values=["Hızlı (200 dpi)", "Normal (300 dpi)", "Yüksek (400 dpi)"], state="readonly")
        self.dpi.current(1)
        self.satir("Kalite:", self.dpi)
        self.sadece = tk.BooleanVar(value=True)
        self.tam_satir(ttk.Checkbutton(self.govde, text="Zaten yazısı olan sayfalara dokunma", variable=self.sadece,
                                       style="Kart.TCheckbutton"))

    def topla(self):
        d = "+".join(k for k, v in (("tur", self.tur), ("eng", self.eng)) if v.get()) or "tur"
        return {"dil": d, "dpi": (200, 300, 400)[self.dpi.current()], "sadece": self.sadece.get()}


# ================================================================== metin penceresi
class MetinPenceresi(tk.Toplevel):
    def __init__(self, ust, metin, kaynak_ad, kaydet_yolu_sor):
        super().__init__(ust)
        self.title(f"Metin — {kaynak_ad}")
        self.geometry("760x620")
        self.configure(bg=R_KART)
        self.kutu = scrolledtext.ScrolledText(self, wrap="word", font=("Consolas" if WINDOWS else "DejaVu Sans Mono", 10),
                                              relief="flat", padx=12, pady=10)
        self.kutu.pack(fill="both", expand=True)
        self.kutu.insert("1.0", metin)
        d = ttk.Frame(self, style="Kart.TFrame", padding=10)
        d.pack(fill="x")

        def kopyala():
            self.clipboard_clear()
            self.clipboard_append(self.kutu.get("1.0", "end-1c"))
            durum.config(text="✔ Panoya kopyalandı")

        def kaydet(tur):
            y = kaydet_yolu_sor(tur)
            if not y:
                return
            m = self.kutu.get("1.0", "end-1c")
            if tur == "txt":
                with open(y, "w", encoding="utf-8-sig") as f:
                    f.write(m)
            else:
                import docx
                belge = docx.Document()
                for p in m.split("\n"):
                    if p.startswith("========== Sayfa"):
                        belge.add_heading(p.strip("= "), level=2)
                    else:
                        belge.add_paragraph(p)
                belge.save(y)
            durum.config(text=f"✔ Kaydedildi: {os.path.basename(y)}")

        ttk.Button(d, text="📋 Tümünü Kopyala", style="Ana.TButton", command=kopyala).pack(side="left")
        ttk.Button(d, text="💾 TXT Kaydet", style="Ikincil.TButton", command=lambda: kaydet("txt")).pack(side="left", padx=6)
        ttk.Button(d, text="💾 Word Kaydet", style="Ikincil.TButton", command=lambda: kaydet("docx")).pack(side="left")
        durum = ttk.Label(d, text="", style="Soluk.TLabel")
        durum.pack(side="right")
