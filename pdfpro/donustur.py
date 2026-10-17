# -*- coding: utf-8 -*-
"""
PDF PRO - Donusturme
PDF -> Word / Excel / Resim / Metin
Resim / Word / Excel -> PDF
"""
import os
import re
import shutil
import subprocess
import sys
import tempfile

from .cekirdek import fitz, pdf_ac, kaydet, PdfProHata, BilesenEksik, sayfa_araligi_coz, benzersiz_yol, RESIM_UZANTILARI

WINDOWS = sys.platform.startswith("win")


# ------------------------------------------------------------------ PDF -> Word
def pdf_to_word(yol, cikti, ilerleme=None):
    try:
        from pdf2docx import Converter
    except ImportError:
        raise PdfProHata("Word dönüştürücü (pdf2docx) kurulu değil.\nKURULUM.bat dosyasını tekrar çalıştırın.")
    with pdf_ac(yol) as d:
        n = d.page_count
        metinsiz = sum(1 for p in d if len(p.get_text().strip()) < 25)
    if ilerleme:
        ilerleme(0, 1, f"{n} sayfa Word'e çevriliyor")
    cv = Converter(yol)
    try:
        cv.convert(cikti, start=0, end=None)
    finally:
        cv.close()
    if ilerleme:
        ilerleme(1, 1, "Bitti")
    uyari = None
    if metinsiz == n:
        uyari = ("Bu PDF taranmış (resim) görünüyor; Word'de yazı düzenlenemez.\n"
                 "Önce 'OCR & Metin' sekmesinden 'Aranabilir PDF' yapıp sonra çevirin\n"
                 "ya da 'Metni Çıkar' ile düz metin alın.")
    return cikti, uyari


# ------------------------------------------------------------------ PDF -> Excel
def _satirlari_sutunla(page):
    """Tablo bulunamazsa: kelimeleri satir ve bosluk mesafesine gore hucrelere boler."""
    kelimeler = page.get_text("words", sort=True)  # x0,y0,x1,y1,word,block,line,wno
    if not kelimeler:
        return []
    kelimeler = sorted(kelimeler, key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    satirlar = []
    mevcut = []
    son_y = None
    for w in kelimeler:
        orta_y = (w[1] + w[3]) / 2
        yuk = w[3] - w[1]
        if son_y is not None and abs(orta_y - son_y) > yuk * 0.6:
            satirlar.append(mevcut)
            mevcut = []
            son_y = None
        mevcut.append(w)
        if son_y is None:
            son_y = orta_y
    if mevcut:
        satirlar.append(mevcut)

    tablo = []
    for s in satirlar:
        s.sort(key=lambda w: w[0])
        hucreler = []
        hucre = s[0][4]
        for onceki, w in zip(s, s[1:]):
            yuk = max(w[3] - w[1], 1)
            bosluk = w[0] - onceki[2]
            if bosluk > yuk * 1.2:   # genis bosluk = yeni sutun
                hucreler.append(hucre)
                hucre = w[4]
            else:
                hucre += " " + w[4]
        hucreler.append(hucre)
        tablo.append(hucreler)
    return tablo


def _sayiya_cevir(deger):
    if deger is None:
        return None
    s = str(deger).strip()
    if not s:
        return ""
    t = s.replace(" ", "")
    isaret = 1
    if t.startswith("-"):
        isaret, t = -1, t[1:]
    if not t or len(t.replace(".", "").replace(",", "")) > 15:
        return s  # telefon / uzun kod: metin kalsin
    if t.isdigit():
        return s if (len(t) > 1 and t.startswith("0")) else isaret * int(t)
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", t):        # 12.500,50 (TR)
        return isaret * float(t.replace(".", "").replace(",", "."))
    if re.fullmatch(r"\d+,\d+", t):                          # 15,25
        return isaret * float(t.replace(",", "."))
    if re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", t):        # 12,500.50 (EN)
        return isaret * float(t.replace(",", ""))
    if re.fullmatch(r"\d+\.\d+", t):                         # 15.25
        return isaret * float(t)
    return s


def pdf_to_excel(yol, cikti, sayilari_cevir=True, ilerleme=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    doc = pdf_ac(yol)
    wb = Workbook()
    wb.remove(wb.active)
    ince = Side(style="thin", color="B0B7C3")
    kenar = Border(left=ince, right=ince, top=ince, bottom=ince)
    baslik_dolgu = PatternFill("solid", fgColor="1F3A5F")
    toplam_tablo = 0
    metinsiz = 0

    for i, page in enumerate(doc):
        if ilerleme:
            ilerleme(i, doc.page_count, f"Sayfa {i + 1} tablo aranıyor")
        tablolar = []
        try:
            bulunan = page.find_tables()
            for t in bulunan.tables:
                veri = t.extract()
                if veri and len(veri) >= 1 and max(len(r) for r in veri) >= 2:
                    tablolar.append(veri)
        except Exception:
            pass
        if not tablolar:
            serbest = _satirlari_sutunla(page)
            if serbest:
                tablolar = [serbest]
            elif not page.get_text().strip():
                metinsiz += 1
        for ti, veri in enumerate(tablolar):
            ad = f"Sayfa{i + 1}" + (f"_T{ti + 1}" if len(tablolar) > 1 else "")
            ws = wb.create_sheet(ad[:31])
            for r, satir in enumerate(veri, start=1):
                for c, hucre in enumerate(satir, start=1):
                    deger = (hucre or "").replace("\n", " ").strip() if isinstance(hucre, str) or hucre is None else hucre
                    if sayilari_cevir and r > 1:
                        deger = _sayiya_cevir(deger)
                    cell = ws.cell(row=r, column=c, value=deger)
                    cell.border = kenar
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if r == 1:
                        cell.font = Font(bold=True, color="FFFFFF")
                        cell.fill = baslik_dolgu
            # sutun genislikleri
            for c in range(1, ws.max_column + 1):
                uz = max((len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, ws.max_row + 1)), default=8)
                ws.column_dimensions[get_column_letter(c)].width = min(max(8, uz + 2), 60)
            ws.freeze_panes = "A2"
            toplam_tablo += 1
    doc.close()
    if toplam_tablo == 0:
        raise PdfProHata("PDF'te okunabilir tablo/metin bulunamadı.\n"
                         "Taranmış bir belgeyse önce 'OCR & Metin' sekmesinden 'Aranabilir PDF' yapın.")
    wb.save(cikti)
    uyari = None
    if metinsiz:
        uyari = f"{metinsiz} sayfa taranmış resim olduğu için atlandı (önce OCR yapın)."
    return cikti, toplam_tablo, uyari


# ------------------------------------------------------------------ PDF -> Resim
def pdf_to_resim(yol, klasor, bicim="png", dpi=200, aralik="", ilerleme=None):
    doc = pdf_ac(yol)
    sayfalar = sayfa_araligi_coz(aralik, doc.page_count)
    ad = os.path.splitext(os.path.basename(yol))[0]
    os.makedirs(klasor, exist_ok=True)
    ciktilar = []
    for k, s in enumerate(sayfalar):
        if ilerleme:
            ilerleme(k, len(sayfalar), f"Sayfa {s + 1}")
        pix = doc[s].get_pixmap(dpi=int(dpi), alpha=False)
        c = benzersiz_yol(os.path.join(klasor, f"{ad}_s{s + 1}.{bicim}"))
        if bicim in ("jpg", "jpeg"):
            pix.save(c, jpg_quality=90)
        else:
            pix.save(c)
        ciktilar.append(c)
    doc.close()
    return ciktilar


# ------------------------------------------------------------------ Resim -> PDF
def resim_to_pdf(resimler, cikti, sayfa_boyutu="orijinal", kenar_mm=0, ilerleme=None):
    """sayfa_boyutu: 'orijinal' | 'A4' | 'A4 yatay'"""
    if not resimler:
        raise PdfProHata("Resim seçilmedi.")
    hedef = fitz.open()
    for i, yol in enumerate(resimler):
        if ilerleme:
            ilerleme(i, len(resimler), os.path.basename(yol))
        img = fitz.open(yol)
        rect = img[0].rect
        pdfb = img.convert_to_pdf()
        img.close()
        tek = fitz.open("pdf", pdfb)
        if sayfa_boyutu == "orijinal":
            hedef.insert_pdf(tek)
        else:
            w, h = fitz.paper_size("a4")
            if sayfa_boyutu == "A4 yatay" or (sayfa_boyutu == "A4 otomatik" and rect.width > rect.height):
                w, h = h, w
            p = hedef.new_page(width=w, height=h)
            k = kenar_mm * 72 / 25.4
            p.show_pdf_page(fitz.Rect(k, k, w - k, h - k), tek, 0, keep_proportion=True)
        tek.close()
    kaydet(hedef, cikti)
    hedef.close()
    return cikti


# ------------------------------------------------------------------ Office -> PDF
def _libreoffice_bul():
    for ad in ("soffice", "libreoffice"):
        y = shutil.which(ad)
        if y:
            return y
    for y in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if os.path.isfile(y):
            return y
    return None


def _libreoffice_cevir(yol, cikti):
    lo = _libreoffice_bul()
    if not lo:
        return False
    tmp = tempfile.mkdtemp(prefix="pdfpro_")
    try:
        subprocess.run([lo, "--headless", "--convert-to", "pdf", "--outdir", tmp, yol],
                       check=True, timeout=300, capture_output=True,
                       creationflags=0x08000000 if WINDOWS else 0)
        uret = os.path.join(tmp, os.path.splitext(os.path.basename(yol))[0] + ".pdf")
        if not os.path.isfile(uret):
            return False
        shutil.move(uret, cikti)
        return True
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _com_baslat():
    import pythoncom
    pythoncom.CoInitialize()


def word_to_pdf(yol, cikti, ilerleme=None):
    yol = os.path.abspath(yol)
    cikti = os.path.abspath(cikti)
    if ilerleme:
        ilerleme(0, 1, "Word açılıyor")
    hata = None
    if WINDOWS:
        try:
            _com_baslat()
            import win32com.client
            word = win32com.client.DispatchEx("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            try:
                belge = word.Documents.Open(yol, ReadOnly=True, ConfirmConversions=False, AddToRecentFiles=False)
                belge.ExportAsFixedFormat(cikti, 17)  # 17 = wdExportFormatPDF
                belge.Close(False)
            finally:
                word.Quit()
            return cikti
        except Exception as e:
            hata = e
    if _libreoffice_cevir(yol, cikti):
        return cikti
    raise PdfProHata("Word → PDF için bilgisayarda Microsoft Word (veya LibreOffice) gerekli.\n"
                     + (f"\nAyrıntı: {hata}" if hata else ""))


def excel_to_pdf(yol, cikti, tek_sayfaya_sigdir=True, ilerleme=None):
    yol = os.path.abspath(yol)
    cikti = os.path.abspath(cikti)
    if ilerleme:
        ilerleme(0, 1, "Excel açılıyor")
    hata = None
    if WINDOWS:
        try:
            _com_baslat()
            import win32com.client
            xl = win32com.client.DispatchEx("Excel.Application")
            xl.Visible = False
            xl.DisplayAlerts = False
            try:
                wb = xl.Workbooks.Open(yol, ReadOnly=True)
                if tek_sayfaya_sigdir:
                    for ws in wb.Worksheets:
                        try:
                            ws.PageSetup.Zoom = False
                            ws.PageSetup.FitToPagesWide = 1
                            ws.PageSetup.FitToPagesTall = False
                        except Exception:
                            pass
                wb.ExportAsFixedFormat(0, cikti)  # 0 = xlTypePDF
                wb.Close(False)
            finally:
                xl.Quit()
            return cikti
        except Exception as e:
            hata = e
    if _libreoffice_cevir(yol, cikti):
        return cikti
    raise PdfProHata("Excel → PDF için bilgisayarda Microsoft Excel (veya LibreOffice) gerekli.\n"
                     + (f"\nAyrıntı: {hata}" if hata else ""))


def herhangi_to_pdf(yol, cikti, ilerleme=None):
    uz = os.path.splitext(yol)[1].lower()
    if uz in (".doc", ".docx", ".rtf", ".odt", ".txt"):
        return word_to_pdf(yol, cikti, ilerleme)
    if uz in (".xls", ".xlsx", ".xlsm", ".csv", ".ods"):
        return excel_to_pdf(yol, cikti, ilerleme=ilerleme)
    if uz in RESIM_UZANTILARI:
        return resim_to_pdf([yol], cikti, ilerleme=ilerleme)
    if uz in (".ppt", ".pptx", ".pps", ".ppsx", ".odp"):
        return pptx_to_pdf(yol, cikti, ilerleme)
    if uz in (".htm", ".html", ".mht", ".mhtml"):
        return html_to_pdf(yol, cikti, ilerleme)
    raise PdfProHata(f"Bu dosya türü desteklenmiyor: {uz}")


# ================================================================== PowerPoint
def pdf_to_pptx(yol, cikti, duzenlenebilir=True, ilerleme=None, iptal=None):
    """Her sayfa bir slayt. duzenlenebilir=True: arka plan (yazisiz) resim + ustunde degistirilebilir
    yazi kutulari. False: sayfa birebir resim olarak (gorunum %100 ayni, yazi duzenlenemez)."""
    try:
        from pptx import Presentation
        from pptx.util import Emu, Pt
        from pptx.dml.color import RGBColor
    except ImportError:
        raise BilesenEksik("PowerPoint dönüştürme bileşeni", "python-pptx")
    import io
    doc = pdf_ac(yol)
    prs = Presentation()
    EMU = 12700                       # 1 pt = 12700 EMU
    ilk = doc[0].rect
    prs.slide_width, prs.slide_height = Emu(int(ilk.width * EMU)), Emu(int(ilk.height * EMU))
    bos = prs.slide_layouts[6]
    n = doc.page_count
    for i in range(n):
        if iptal and iptal():
            raise PdfProHata("İşlem iptal edildi.")
        if ilerleme:
            ilerleme(i, n, f"Slayt {i + 1}/{n}")
        page = doc[i]
        r = page.rect
        sl = prs.slides.add_slide(bos)
        olc = (prs.slide_width / EMU) / r.width
        satirlar = []
        if duzenlenebilir:
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") != 0:
                    continue
                for ln in b["lines"]:
                    sp = [s for s in ln["spans"] if s["text"].strip() and s.get("alpha", 255) != 0]
                    if sp and abs(ln["dir"][1]) < 0.01:
                        satirlar.append((fitz.Rect(ln["bbox"]), sp))
        # arka plan: yazilar silinmis sayfa resmi
        tmp = fitz.open()
        tmp.insert_pdf(doc, from_page=i, to_page=i)
        tp = tmp[0]
        if satirlar:
            from .duzenle import icerik_normallestir, sayfa_hazirla
            sayfa_hazirla(tp)
            icerik_normallestir(tp)      # resmi belgelerdeki ' operatoru silmede yazilari kaydirmasin
            tp = tmp[0]
            for rr, _ in satirlar:
                tp.add_redact_annot(rr, fill=None, cross_out=False)
            try:
                tp.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,
                                    graphics=getattr(fitz, "PDF_REDACT_LINE_ART_NONE", 0))
            except TypeError:
                tp.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)
        pix = tp.get_pixmap(dpi=150, alpha=False)
        tmp.close()
        sl.shapes.add_picture(io.BytesIO(pix.tobytes("png")), 0, 0, prs.slide_width, int(r.height * olc * EMU))
        for rr, sp in satirlar:
            kutu = sl.shapes.add_textbox(Emu(int(rr.x0 * olc * EMU)), Emu(int((rr.y0 - 1) * olc * EMU)),
                                         Emu(int((rr.width + 4) * olc * EMU)), Emu(int((rr.height + 2) * olc * EMU)))
            tf = kutu.text_frame
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            tf.word_wrap = False
            p = tf.paragraphs[0]
            for s in sp:
                run = p.add_run()
                run.text = s["text"]
                f = run.font
                f.size = Pt(max(4, s["size"] * olc))
                fname = s["font"].split("+")[-1]
                f.name = fname.split(",")[0].replace("-", " ").replace("Bold", "").replace("Italic", "").strip() or "Arial"
                f.bold = bool(s["flags"] & 16) or "bold" in s["font"].lower()
                f.italic = bool(s["flags"] & 2) or "italic" in s["font"].lower()
                c = s["color"]
                f.color.rgb = RGBColor((c >> 16) & 255, (c >> 8) & 255, c & 255)
    doc.close()
    prs.save(cikti)
    return cikti


def pptx_to_pdf(yol, cikti, ilerleme=None):
    yol, cikti = os.path.abspath(yol), os.path.abspath(cikti)
    if ilerleme:
        ilerleme(0, 1, "PowerPoint açılıyor")
    hata = None
    if WINDOWS:
        try:
            _com_baslat()
            import win32com.client
            pp = win32com.client.DispatchEx("PowerPoint.Application")
            try:
                sunum = pp.Presentations.Open(yol, ReadOnly=True, Untitled=False, WithWindow=False)
                sunum.SaveAs(cikti, 32)        # 32 = ppSaveAsPDF
                sunum.Close()
            finally:
                pp.Quit()
            return cikti
        except Exception as e:
            hata = e
    if _libreoffice_cevir(yol, cikti):
        return cikti
    raise PdfProHata("PowerPoint → PDF için bilgisayarda Microsoft PowerPoint (veya LibreOffice) gerekli.\n"
                     + (f"\nAyrıntı: {hata}" if hata else ""))


# ================================================================== HTML / web sayfasi
def _tarayici_bul():
    adaylar = []
    if WINDOWS:
        for kok in (os.environ.get("PROGRAMFILES(X86)", r"C:\Program Files (x86)"),
                    os.environ.get("PROGRAMFILES", r"C:\Program Files"),
                    os.environ.get("LOCALAPPDATA", "")):
            adaylar += [os.path.join(kok, r"Microsoft\Edge\Application\msedge.exe"),
                        os.path.join(kok, r"Google\Chrome\Application\chrome.exe")]
    adaylar += [os.environ.get("PDFPRO_TARAYICI", ""), "/opt/pw-browsers/chromium"]
    for ad in ("msedge", "chrome", "chromium", "chromium-browser", "google-chrome"):
        y = shutil.which(ad)
        if y:
            adaylar.append(y)
    for a in adaylar:
        if a and os.path.isfile(a):
            return a
    return None


def html_to_pdf(kaynak, cikti, ilerleme=None):
    """kaynak: web adresi (https://...) veya .html dosyasi. Edge / Chrome ile yazdirir."""
    tarayici = _tarayici_bul()
    if not tarayici:
        raise PdfProHata("HTML → PDF için Microsoft Edge veya Google Chrome bulunamadı.")
    if os.path.isfile(kaynak):
        from pathlib import Path
        adres = Path(os.path.abspath(kaynak)).as_uri()
    else:
        adres = kaynak if re.match(r"^[a-z]+://", kaynak, re.I) else "https://" + kaynak
    if ilerleme:
        ilerleme(0, 1, "Sayfa yükleniyor")
    cikti = os.path.abspath(cikti)
    profil = tempfile.mkdtemp(prefix="pdfpro_tar_")
    komut = [tarayici, "--headless=new", "--disable-gpu", "--no-sandbox", "--no-first-run",
             f"--user-data-dir={profil}", "--no-pdf-header-footer", "--run-all-compositor-stages-before-draw",
             "--virtual-time-budget=8000", f"--print-to-pdf={cikti}", adres]
    try:
        subprocess.run(komut, capture_output=True, timeout=120, creationflags=0x08000000 if WINDOWS else 0)
    except subprocess.TimeoutExpired:
        raise PdfProHata("Sayfa 2 dakikada yüklenemedi.")
    finally:
        shutil.rmtree(profil, ignore_errors=True)
    if not os.path.isfile(cikti) or os.path.getsize(cikti) < 200:
        raise PdfProHata(f"Sayfa PDF'e çevrilemedi:\n{adres}")
    return cikti


# ================================================================== Markdown
def pdf_to_markdown(yol, cikti=None, ilerleme=None):
    """Basliklar (yazi boyuna gore), listeler ve tablolarla Markdown metni uretir."""
    doc = pdf_ac(yol)
    boylar = {}
    for p in doc:
        for b in p.get_text("dict")["blocks"]:
            for ln in b.get("lines", []):
                for s in ln["spans"]:
                    if s["text"].strip():
                        k = round(s["size"])
                        boylar[k] = boylar.get(k, 0) + len(s["text"])
    govde = max(boylar, key=boylar.get) if boylar else 11
    parcalar = []
    for i, p in enumerate(doc):
        if ilerleme:
            ilerleme(i, doc.page_count, f"Sayfa {i + 1}")
        tablo_alanlari = []
        tablolar = []
        try:
            for t in p.find_tables().tables:
                veri = t.extract()
                if veri and len(veri) >= 2 and max(len(r) for r in veri) >= 2:
                    tablo_alanlari.append(fitz.Rect(t.bbox))
                    tablolar.append((t.bbox[1], veri))
        except Exception:
            pass
        ogeler = []
        for b in p.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            br = fitz.Rect(b["bbox"])
            if any(br.intersects(t) for t in tablo_alanlari):
                continue
            satirlar = []
            enbuyuk, kalin = 0, False
            for ln in b["lines"]:
                t = "".join(s["text"] for s in ln["spans"]).strip()
                if t:
                    satirlar.append(t)
                    for s in ln["spans"]:
                        enbuyuk = max(enbuyuk, s["size"])
                        kalin = kalin or bool(s["flags"] & 16)
            if not satirlar:
                continue
            metin = " ".join(satirlar)
            if enbuyuk >= govde * 1.6:
                metin = "# " + metin
            elif enbuyuk >= govde * 1.25 or (kalin and len(metin) < 80 and len(satirlar) == 1):
                metin = "## " + metin
            elif re.match(r"^[•\-–▪●]\s*", metin):
                metin = "- " + re.sub(r"^[•\-–▪●]\s*", "", metin)
            ogeler.append((br.y0, metin))
        for y, veri in tablolar:
            satir = lambda r: "| " + " | ".join((c or "").replace("\n", " ").replace("|", "/") for c in r) + " |"
            md = [satir(veri[0]), "|" + "---|" * len(veri[0])] + [satir(r) for r in veri[1:]]
            ogeler.append((y, "\n".join(md)))
        ogeler.sort(key=lambda o: o[0])
        parcalar.append(f"<!-- Sayfa {i + 1} -->\n\n" + "\n\n".join(m for _, m in ogeler))
    doc.close()
    md = "\n\n---\n\n".join(parcalar)
    if cikti:
        with open(cikti, "w", encoding="utf-8") as f:
            f.write(md)
        return cikti
    return md


# ================================================================== PDF/A ve onarim
def _ghostscript_bul():
    for ad in ("gswin64c", "gswin32c", "gs"):
        y = shutil.which(ad)
        if y:
            return y
    for kok in (r"C:\Program Files\gs", r"C:\Program Files (x86)\gs"):
        if os.path.isdir(kok):
            for alt in sorted(os.listdir(kok), reverse=True):
                for ad in ("gswin64c.exe", "gswin32c.exe"):
                    y = os.path.join(kok, alt, "bin", ad)
                    if os.path.isfile(y):
                        return y
    return None


def pdf_a(yol, cikti, ilerleme=None):
    """Uzun sureli arsiv kopyasi (PDF/A-2b). Ghostscript varsa onunla tam uyumlu donusturur;
    yoksa en iyi cabayla (renk profili + arsiv bilgileri) hazirlar ve eksikleri raporlar.
    Donus: (cikti, notlar listesi)"""
    notlar = []
    gs = _ghostscript_bul()
    if gs:
        if ilerleme:
            ilerleme(0, 1, "Ghostscript ile PDF/A")
        r = subprocess.run([gs, "-dPDFA=2", "-dBATCH", "-dNOPAUSE", "-dNOOUTERSAVE", "-sColorConversionStrategy=RGB",
                            "-sDEVICE=pdfwrite", "-dPDFACompatibilityPolicy=1", f"-sOutputFile={cikti}", yol],
                           capture_output=True, timeout=600, creationflags=0x08000000 if WINDOWS else 0)
        if r.returncode == 0 and os.path.isfile(cikti):
            return cikti, ["Ghostscript ile PDF/A-2b olarak dönüştürüldü."]
        notlar.append("Ghostscript dönüştüremedi, yerleşik yöntem kullanıldı.")
    from PIL import ImageCms
    import datetime
    doc = pdf_ac(yol)
    if ilerleme:
        ilerleme(0, 1, "Arşiv bilgileri ekleniyor")
    # renk profili (sRGB) -> OutputIntent
    icc = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes()
    ix = doc.get_new_xref()
    doc.update_object(ix, "<</N 3>>")
    doc.update_stream(ix, icc)
    oi = doc.get_new_xref()
    doc.update_object(oi, f"<</Type/OutputIntent/S/GTS_PDFA1/OutputConditionIdentifier(sRGB IEC61966-2.1)"
                          f"/Info(sRGB IEC61966-2.1)/DestOutputProfile {ix} 0 R>>")
    doc.xref_set_key(doc.pdf_catalog(), "OutputIntents", f"[{oi} 0 R]")
    simdi = datetime.datetime.now().strftime("%Y-%m-%dT%H:%M:%S")
    baslik = (doc.metadata or {}).get("title") or os.path.splitext(os.path.basename(yol))[0]
    xmp = f"""<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>
<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">
<rdf:Description rdf:about="" xmlns:pdfaid="http://www.aiim.org/pdfa/ns/id/" xmlns:dc="http://purl.org/dc/elements/1.1/"
 xmlns:xmp="http://ns.adobe.com/xap/1.0/" xmlns:pdf="http://ns.adobe.com/pdf/1.3/">
<pdfaid:part>2</pdfaid:part><pdfaid:conformance>B</pdfaid:conformance>
<dc:title><rdf:Alt><rdf:li xml:lang="x-default">{baslik}</rdf:li></rdf:Alt></dc:title>
<xmp:CreateDate>{simdi}</xmp:CreateDate><xmp:ModifyDate>{simdi}</xmp:ModifyDate>
<xmp:CreatorTool>PDF PRO</xmp:CreatorTool><pdf:Producer>PDF PRO</pdf:Producer>
</rdf:Description></rdf:RDF></x:xmpmeta>
<?xpacket end="w"?>"""
    doc.set_xml_metadata(xmp)
    doc.set_metadata({"title": baslik, "producer": "PDF PRO", "creator": "PDF PRO"})
    # uyumsuzluklari kontrol et
    gomulmemis = set()
    for p in doc:
        for f in p.get_fonts(full=True):
            if f[1] in ("n/a", ""):
                gomulmemis.add(f[3])
    if gomulmemis:
        notlar.append("Gömülü olmayan yazı tipleri var (" + ", ".join(sorted(gomulmemis)[:5]) +
                      "). Tam uyum için bilgisayara Ghostscript kurulabilir.")
    if doc.is_encrypted:
        notlar.append("Şifreli belgeler PDF/A olamaz; önce şifreyi kaldırın.")
    doc.save(cikti, garbage=3, deflate=True, encryption=fitz.PDF_ENCRYPT_NONE)
    doc.close()
    if not notlar:
        notlar.append("PDF/A-2b arşiv bilgileri ve sRGB renk profili eklendi.")
    return cikti, notlar


def onar(yol, cikti, ilerleme=None):
    """Bozuk / yarim inmis PDF'i kurtarir. Donus: (cikti, kurtarilan_sayfa, rapor)"""
    if ilerleme:
        ilerleme(0, 1, "Dosya inceleniyor")
    try:
        with open(yol, "rb") as f:
            veri = f.read()
    except Exception as e:
        raise PdfProHata(f"Dosya okunamadı:\n{e}")
    rapor = []
    try:
        doc = fitz.open(stream=veri, filetype="pdf")
    except Exception as e:
        # %PDF basligi bozuksa bulup oradan dene
        i = veri.find(b"%PDF")
        if i <= 0:
            raise PdfProHata(f"Bu dosya kurtarılamadı (PDF yapısı bulunamadı).\n{e}")
        doc = fitz.open(stream=veri[i:], filetype="pdf")
        rapor.append("Dosyanın başındaki bozuk kısım atlandı.")
    if doc.needs_pass:
        raise PdfProHata("Dosya şifreli; önce şifreyi kaldırın.")
    if getattr(doc, "is_repaired", False):
        rapor.append("Dosyanın iç yapısı (xref tablosu) bozuktu, yeniden kuruldu.")
    yeni = fitz.open()
    kurtarilan = 0
    gorulen = set()
    import hashlib
    for i in range(doc.page_count):
        try:
            pix = doc[i].get_pixmap(matrix=fitz.Matrix(0.15, 0.15), alpha=False)   # cizilebiliyor mu
            ozet = hashlib.md5(pix.samples).hexdigest()
            if ozet in gorulen:
                continue            # bozuk dosyada ayni sayfanin eski kopyalari: atla
            gorulen.add(ozet)
            yeni.insert_pdf(doc, from_page=i, to_page=i)
            kurtarilan += 1
        except Exception:
            rapor.append(f"Sayfa {i + 1} kurtarılamadı.")
    if kurtarilan == 0:
        raise PdfProHata("Hiçbir sayfa kurtarılamadı.")
    yeni.save(cikti, garbage=4, deflate=True)
    yeni.close()
    doc.close()
    if not rapor:
        rapor.append("Dosyada önemli bir bozukluk bulunmadı; temiz bir kopya yeniden yazıldı.")
    return cikti, kurtarilan, rapor
